"""Offline regressietests: python -m unittest discover -s tests -v."""
import ast
import copy
import importlib
import io
import json
from pathlib import Path
import sys
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch
from contextlib import ExitStack

import docx
import pandas as pd
from streamlit.testing.v1 import AppTest
from streamlit.runtime.scriptrunner import StopException

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def toets():
    return {"vragen": [
        {"id": i, "type": "mc", "vraag": f"Vraag {i}",
         "opties": ["A) Een", "B) Twee", "C) Drie", "D) Vier"], "correct": "A"}
        for i in range(1, 5)
    ] + [{"id": i, "type": "open", "vraag": f"Vraag {i}"} for i in (5, 6)]}


class Regressies(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Gebruik de echte configuratieconstanten, zonder clients of secrets te laden.
        cfg = ModuleType("config")
        tree = ast.parse((ROOT / "config.py").read_text(encoding="utf-8"))
        names = {"NIVEAUS", "LEERJAREN_CLUSTERS", "HOOFDSTUKKEN", "VERSIE", "LAATSTE_UPDATE"}
        for node in tree.body:
            if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name) and node.targets[0].id in names:
                setattr(cfg, node.targets[0].id, ast.literal_eval(node.value))
        cfg.ALLE_CLUSTERS = [k for ks in cfg.NIVEAUS.values() for k in ks]
        cfg.supabase = MagicMock()
        cfg.ai_client = MagicMock()
        cfg.setup_page = lambda: None
        cfg.pas_styling_toe = lambda: None
        cls.config_patch = patch.dict(sys.modules, {"config": cfg})
        cls.config_patch.start()
        cls.bestanden = importlib.import_module("bestanden")
        cls.auth = importlib.import_module("auth")
        cls.ai = importlib.import_module("ai_docent")
        cls.ui = importlib.import_module("ui_leerling")
        cls.beheer = importlib.import_module("ui_beheer")

    @classmethod
    def tearDownClass(cls):
        cls.config_patch.stop()

    def test_import_leerjaar(self):
        self.assertEqual(self.ui.get_leerjaar("4Hak1"), "4Havo")

    def test_paginering_ook_bij_lagere_serverlimiet(self):
        rows = [{"PogingID": str(i)} for i in range(1205)]
        query = MagicMock()
        query.select.return_value = query
        query.order.return_value = query
        query.range.side_effect = lambda lo, hi: SimpleNamespace(
            execute=lambda: SimpleNamespace(data=rows[lo:min(hi + 1, lo + 200)]))
        with patch.object(self.bestanden, "supabase") as db:
            db.table.return_value = query
            self.assertEqual(self.bestanden.haal_tabel_op("resultaten", "PogingID"), rows)
        self.assertEqual(query.range.call_count, 8)

    def test_geen_ongewijzigde_accounts_of_velden_overschrijven(self):
        db = {"a": {"Gebruikersnaam": "a", "Goedgekeurd": "Nee", "WachtwoordHash": "oud"},
              "b": {"Gebruikersnaam": "b", "Goedgekeurd": "Nee"}}
        eerste = self.auth.AccountOverzicht(copy.deepcopy(db))
        tweede = self.auth.AccountOverzicht(copy.deepcopy(db))
        eerste["a"]["Goedgekeurd"] = "Ja"
        tweede["b"]["Goedgekeurd"] = "Ja"
        tweede["a"]["WachtwoordHash"] = "nieuw"
        query = MagicMock()
        def update(fields):
            def eq(key, account):
                def execute():
                    db[account].update(fields)
                    return SimpleNamespace(data=[db[account]])
                return SimpleNamespace(execute=execute)
            return SimpleNamespace(eq=eq)
        query.update.side_effect = update
        with patch.object(self.auth, "supabase") as client:
            client.table.return_value = query
            self.auth.bewaar_alle_gebruikers(eerste)
            self.auth.bewaar_alle_gebruikers(tweede)
        self.assertEqual(db["a"]["Goedgekeurd"], "Ja")
        self.assertEqual(db["b"]["Goedgekeurd"], "Ja")
        self.assertEqual(db["a"]["WachtwoordHash"], "nieuw")
        query.upsert.assert_not_called()

    def test_nieuw_account_gebruikt_insert(self):
        accounts = self.auth.AccountOverzicht({})
        accounts["a"] = {"Gebruikersnaam": "a"}
        with patch.object(self.auth, "supabase") as db:
            db.table.return_value.insert.return_value.execute.return_value.data = [accounts["a"]]
            self.auth.bewaar_alle_gebruikers(accounts)
            db.table.return_value.insert.assert_called_once_with(accounts["a"])
            db.table.return_value.upsert.assert_not_called()

    def test_opslaan_stopt_bij_databasefout(self):
        accounts = self.auth.AccountOverzicht({})
        accounts["a"] = {"Gebruikersnaam": "a"}
        with patch.object(self.auth, "supabase") as db, patch.object(self.auth, "st") as st:
            st.stop.side_effect = StopException
            db.table.return_value.insert.return_value.execute.side_effect = RuntimeError("offline")
            with self.assertRaises(StopException):
                self.auth.bewaar_alle_gebruikers(accounts)
            st.error.assert_called_once()

    def test_verwijderen_stopt_bij_geen_verwijderde_rijen(self):
        with patch.object(self.auth, "supabase") as db, patch.object(self.auth, "st") as st:
            st.stop.side_effect = StopException
            db.table.return_value.delete.return_value.eq.return_value.execute.return_value.data = []
            with self.assertRaises(StopException):
                self.auth.verwijder_account("gebruikers", "Gebruikersnaam", "a")

    def test_leesfout_is_geen_leeg_accountbestand(self):
        with patch.object(self.auth, "haal_tabel_op", side_effect=RuntimeError("offline")), patch.object(self.auth, "st") as st:
            st.stop.side_effect = StopException
            with self.assertRaises(StopException):
                self.auth.laad_gebruikers()

    def test_docent_wijzigingen_serialiseren_alleen_gewijzigde_klassen(self):
        docs = self.auth.AccountOverzicht({"d": {"DocentID": "d", "Klassen": ["4Hak1"], "Naam": "Docent"}})
        docs["d"]["Klassen"].append("4Hak2")
        with patch.object(self.auth, "supabase") as db:
            db.table.return_value.update.return_value.eq.return_value.execute.return_value.data = [{}]
            self.auth.bewaar_alle_docenten(docs)
            db.table.return_value.update.assert_called_once_with({"Klassen": "4Hak1,4Hak2"})

    def test_ongeldige_toets_wordt_afgewezen(self):
        cases = []
        missing = toets(); del missing["vragen"][0]["correct"]; cases.append(missing)
        duplicate = toets(); duplicate["vragen"][1]["id"] = 1; cases.append(duplicate)
        short = toets(); short["vragen"].pop(); cases.append(short)
        for data in cases:
            with self.subTest(data=data), self.assertRaises(ValueError):
                self.ai.valideer_toets(data)

    def test_score_grenzen_en_meerkeuze(self):
        answers = {i: "A) Een" for i in range(1, 5)}
        for score in [-1, 7, float("nan"), float("inf"), None, True, "verkeerd", 6]:
            evaluation = {"score_open_vragen_totaal": score,
                          "beoordeling_open_vragen": "Feedback", "docenten_feedback": "Goed gewerkt"}
            with self.subTest(score=score), patch.object(self.ai, "ai_client") as client:
                client.models.generate_content.return_value.text = json.dumps(evaluation)
                if score == 6:
                    self.assertEqual(self.ai.kijk_toets_na("Havo", "A", "Theorie", toets(), answers)[0], 10)
                else:
                    with self.assertRaises((ValueError, TypeError)):
                        self.ai.kijk_toets_na("Havo", "A", "Theorie", toets(), answers)

    def test_docx_tabellen_en_volgorde(self):
        document = docx.Document()
        document.add_paragraph("Voor")
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Tabel"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "Genest"
        document.add_paragraph("Na")
        content = io.BytesIO(); document.save(content)
        with patch.object(self.bestanden, "supabase") as db:
            db.storage.from_.return_value.download.return_value = content.getvalue()
            result = self.bestanden.lees_docx("4Havo", "H4H1", "a.docx")
        self.assertEqual([x for x in result.splitlines() if x], ["Voor", "Tabel", "Genest", "Na"])

    def test_gelijknamige_lessen_en_oude_resultaten(self):
        lessons = pd.Series(["4Havo/H4H1/a.docx", "4Havo/H4H2/a.docx", "a.docx"])
        with patch.object(self.bestanden, "haal_bestanden_op", return_value=["a.docx"]), patch.object(self.bestanden.st, "warning"):
            self.assertEqual(self.bestanden.les_resultaat_mask(lessons, "4Havo", "H4H1", "a.docx").tolist(), [True, False, False])
        with patch.object(self.bestanden, "haal_bestanden_op", side_effect=lambda year, chapter: ["a.docx"] if chapter == "H4H1" else []):
            self.assertEqual(self.bestanden.les_resultaat_mask(lessons, "4Havo", "H4H1", "a.docx").tolist(), [True, False, True])

    def leerling_app(self, rows):
        stack = ExitStack(); self.addCleanup(stack.close)
        stack.enter_context(patch.object(self.ui, "haal_alle_resultaten_op", side_effect=lambda: pd.DataFrame(rows)))
        stack.enter_context(patch.object(self.ui, "haal_bestanden_op", return_value=["a.docx"]))
        stack.enter_context(patch.object(self.ui, "lees_docx", side_effect=lambda year, chapter, file: chapter))
        generator = stack.enter_context(patch.object(self.ui, "genereer_toets_gecached", side_effect=lambda *args: toets()))
        stack.enter_context(patch.object(self.ui, "kijk_toets_na", return_value=(8.0, "Feedback", "Goed gewerkt")))
        def save(*args):
            rows.append({"PogingID": str(len(rows)), "Gebruikersnaam": "a", "Cluster": "4Hak1",
                         "Les": args[5], "Cijfer": args[6], "Beoordeling": args[7],
                         "Tijdstip": "2026-10-02", "ReactieGelezen": "True", "DocentReactie": ""})
            return True, ""
        stack.enter_context(patch.object(self.ui, "sla_resultaat_op", side_effect=save))
        app = AppTest.from_file(str(ROOT / "main.py"), default_timeout=10)
        for key, value in dict(ingelogd=True, rol="leerling", gebruikersnaam="a", voornaam="Gaston", cluster="4Hak1", niveau="Havo", is_gast=False).items():
            app.session_state[key] = value
        app.run()
        self.assertFalse(app.exception)
        return app, generator

    def test_feedback_blijft_na_alle_drie_pogingen_en_gaston_is_geen_gast(self):
        rows = []
        app, generator = self.leerling_app(rows)
        app.selectbox(key="ll_kies_les").select("a.docx").run()
        for attempt in range(1, 4):
            next(b for b in app.button if b.label == "Lever in").click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(rows), attempt)
            self.assertEqual(app.session_state["huidig_cijfer"], 8)
            self.assertTrue(any("Eindcijfer: 8.0" in m.value for m in app.markdown))
            next(b for b in app.button if "Terug /" in b.label).click().run()
        self.assertTrue(any("maximale aantal" in e.value for e in app.error))
        self.assertEqual(generator.call_count, 3)

    def test_hoofdstukwissel_vernieuwt_vragen_en_antwoorden(self):
        app, generator = self.leerling_app([])
        app.selectbox(key="ll_kies_les").select("a.docx").run()
        app.text_area[0].set_value("Oud antwoord").run()
        app.selectbox(key="ll_kies_hst").select("H4H2").run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state["huidige_les"], "4Havo/H4H2/a.docx")
        self.assertEqual(app.text_area[0].value, "")
        self.assertEqual(generator.call_count, 2)

    def test_feedbackmelding_ververst(self):
        rows = []
        app, _ = self.leerling_app(rows)
        rows.append({"PogingID": "1", "Gebruikersnaam": "a", "Cluster": "4Hak1", "Les": "oud",
                     "Cijfer": 8, "Beoordeling": "Goed", "Tijdstip": "2026-10-02",
                     "DocentReactie": "Reactie", "ReactieGelezen": "False"})
        app.run()
        self.assertTrue(any("Nieuw bericht" in e.value for e in app.error))
        rows[0]["ReactieGelezen"] = "True"
        app.run()
        self.assertFalse(any("Nieuw bericht" in e.value for e in app.error))

    def test_main_start_en_lege_admin_secret(self):
        with patch.object(self.auth, "laad_docenten", return_value={}):
            app = AppTest.from_file(str(ROOT / "main.py"), default_timeout=10)
            app.secrets["TEST"] = "offline"
            app.run()
            self.assertFalse(app.exception)
            app.text_input(key="d_login").set_value("admin")
            next(b for b in app.button if b.label == "Log in als docent").click().run()
            self.assertFalse(app.exception)
            self.assertNotIn("ingelogd", app.session_state)

    def test_opslagfout_toont_geen_ingeleverd_en_behoudt_antwoorden(self):
        app, _ = self.leerling_app([])
        app.selectbox(key="ll_kies_les").select("a.docx").run()
        app.text_area[0].set_value("Mijn antwoord")
        with patch.object(self.ui, "sla_resultaat_op", return_value=(False, "offline")):
            next(b for b in app.button if b.label == "Lever in").click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.session_state["toets_ingeleverd"])
        self.assertEqual(app.text_area[0].value, "Mijn antwoord")
        self.assertTrue(any("database weigerde" in e.value for e in app.error))

    def test_echte_gast_behoudt_gastfuncties(self):
        app, _ = self.leerling_app([])
        app.session_state["is_gast"] = True
        app.session_state["voornaam"] = "Bezoeker"
        app.run()
        self.assertFalse(app.exception)
        self.assertTrue(any("Gasten hebben geen instellingen" in e.value for e in app.warning))

    def test_admin_niveau_ververst_klassen_voor_verzenden(self):
        with patch.object(self.beheer, "laad_gebruikers", side_effect=lambda: self.auth.AccountOverzicht({})), patch.object(self.beheer, "laad_docenten", side_effect=lambda: self.auth.AccountOverzicht({})):
            app = AppTest.from_file(str(ROOT / "main.py"), default_timeout=10)
            for key, value in dict(ingelogd=True, rol="admin", docent_naam="Beheerder").items():
                app.session_state[key] = value
            app.run()
            app.selectbox(key="admin_nieuw_niveau").select("VWO").run()
            self.assertFalse(app.exception)
            klas = next(s for s in app.selectbox if s.label == "Klas:")
            self.assertEqual(klas.options, sys.modules["config"].NIVEAUS["VWO"])

    def test_lesmateriaal_paginering(self):
        files = [{"name": f"les{i}.docx"} for i in range(205)]
        with patch.object(self.bestanden, "supabase") as db:
            db.storage.from_.return_value.list.side_effect = lambda path, options: files[options["offset"]:options["offset"] + options["limit"]]
            self.assertEqual(len(self.bestanden.haal_bestanden_op("4Havo", "H4H1")), 205)


if __name__ == "__main__":
    unittest.main()
